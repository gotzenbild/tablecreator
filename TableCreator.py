import docx
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table, TableRow, TableCell
from odf.text import P


class TableCreator:
    @staticmethod
    def save_table(filename, form, datatable):
        cols = datatable.columnCount()
        rows = datatable.rowCount()
        if form == "Text Files (*.odt)":
            doc = OpenDocumentSpreadsheet()
            table = Table()
            for row in range(rows):
                tr = TableRow()
                for col in range(cols):
                    tc = TableCell(valuetype='string')
                    data = datatable.model().index(row, col).data()
                    if data is None:
                        data = ""
                    tc.addElement(P(text=data))
                    tr.addElement(tc)
                table.addElement(tr)
            doc.spreadsheet.addElement(table)
            doc.save(filename, True)
        elif form == "Text Files (*.docx)":
            doc = docx.Document()
            table = doc.add_table(rows=rows, cols=cols)
            table.style = 'Table Grid'
            for row in range(rows):
                for col in range(cols):
                    cell = table.cell(row, col)
                    data = datatable.model().index(row, col).data()
                    if data is None:
                        data = ""
                    cell.text = data
            doc.save(filename)
